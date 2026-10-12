# -*- coding: utf-8 -*-
"""Общие функции сборки отчётов по практическим работам СИАОД (вариант 14)."""
from copy import deepcopy
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

L = WD_ALIGN_PARAGRAPH.LEFT
C = WD_ALIGN_PARAGRAPH.CENTER
J = WD_ALIGN_PARAGRAPH.JUSTIFY
R = WD_ALIGN_PARAGRAPH.RIGHT


class Report:
    def __init__(self, work_no, theme, part='ч.2'):
        self.doc = Document()
        self._styles()
        self._title(work_no, theme, part)

    # ------------------------------------------------------------------ стили
    def _styles(self):
        doc = self.doc
        st = doc.styles['Normal']
        st.font.name = 'Times New Roman'
        st.font.size = Pt(14)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.first_line_indent = Cm(1.25)
        st.paragraph_format.alignment = J

        sec0 = doc.sections[0]
        sec0.top_margin = sec0.bottom_margin = Cm(2)
        sec0.left_margin = sec0.right_margin = Cm(2)

        for name, size in (('Heading 1', 16), ('Heading 2', 14), ('Heading 3', 14)):
            h = doc.styles[name]
            h.font.name = 'Times New Roman'
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.paragraph_format.line_spacing = 1.5
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.first_line_indent = Cm(0)
            h.paragraph_format.alignment = L

        cs = doc.styles.add_style('CodeBlock', 1)
        cs.base_style = doc.styles['Normal']
        cs.font.name = 'Consolas'
        cs.font.size = Pt(9)
        cs.element.rPr.rFonts.set(qn('w:ascii'), 'Consolas')
        cs.element.rPr.rFonts.set(qn('w:hAnsi'), 'Consolas')
        cs.paragraph_format.line_spacing = 1.0
        cs.paragraph_format.first_line_indent = Cm(0)
        cs.paragraph_format.alignment = L
        cs.paragraph_format.space_after = Pt(0)

    # --------------------------------------------------------- титульный лист
    def _title(self, work_no, theme, part):
        def T(text='', size=14, bold=False, italic=False, align=C,
              first_line=0.0, left=0.0):
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.alignment = align
            pf.line_spacing = 1.0
            pf.space_after = Pt(0)
            pf.first_line_indent = Cm(first_line)
            pf.left_indent = Cm(left)
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
            r.font.name = 'Times New Roman'
            r.font.size = Pt(size)
            r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
            return p

        def gap(n=1, size=12):
            for _ in range(n):
                T(size=size)

        T('МИНОБРНАУКИ РОССИИ')
        T('Федеральное государственное бюджетное образовательное учреждение')
        T('высшего образования')
        T('«МИРЭА – Российский технологический университет»', bold=True)
        T('РТУ МИРЭА', bold=True)
        T('Институт Информационных технологий')
        gap()
        T('Кафедра Математического обеспечения и стандартизации '
          'информационных технологий')
        gap(4)
        T(f'ОТЧЕТ ПО ПРАКТИЧЕСКОЙ РАБОТЕ № {work_no}', bold=True)
        T('по дисциплине', bold=True)
        T(f'«Структуры и алгоритмы обработки данных ({part})»', bold=True)
        gap()
        T(f'Тема: «{theme}»')
        gap(4)
        T('Выполнил студент группы ИКБО-41-25', size=12, align=L, first_line=1.3)
        gap(2)
        T('Мирзоян А.О.', size=12, align=L, first_line=4.16)
        gap()
        T('Принял преподаватель', size=12, align=L, first_line=1.3)
        gap(2)
        T('Скворцова Л.А.', size=12, bold=True, first_line=3.66)
        gap()
        T('Самостоятельная работа выполнена', align=L, first_line=-0.2)
        T('«__»_______202__ г.')
        gap()
        T('(подпись студента)', italic=True, first_line=0.64)
        gap(2)
        T('«Зачтено»', align=L, first_line=-0.2)
        T('«__»_______202__ г')
        gap()
        T('(подпись руководителя)', italic=True)
        gap()
        T('Москва 2026', size=12, align=L, left=6.24, first_line=1.25)

        # титульный лист остаётся в первой секции (поля 2 см),
        # основной текст выносится в новую секцию с полями по ГОСТ
        sec1 = self.doc.add_section(WD_SECTION.NEW_PAGE)
        sec1.top_margin, sec1.bottom_margin = Cm(2), Cm(2)
        sec1.left_margin, sec1.right_margin = Cm(3), Cm(1.5)

    # ------------------------------------------------------------- оглавление
    def toc(self):
        h = self.doc.add_heading('ОГЛАВЛЕНИЕ', level=1)
        h.alignment = C
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        fld = OxmlElement('w:fldSimple')
        fld.set(qn('w:instr'), r'TOC \o "1-2" \h \z \u')
        run = OxmlElement('w:r')
        t = OxmlElement('w:t')
        t.text = 'Для обновления оглавления выделите его и нажмите F9.'
        run.append(t)
        fld.append(run)
        p._p.append(fld)
        self.page_break()

    # ------------------------------------------------------------- содержимое
    def h1(self, text):
        return self.doc.add_heading(text, level=1)

    def h2(self, text):
        return self.doc.add_heading(text, level=2)

    def p(self, text='', bold=False, italic=False, align=None, indent=None,
          size=None, keep=False):
        par = self.doc.add_paragraph()
        r = par.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if align is not None:
            par.paragraph_format.alignment = align
        if indent is not None:
            par.paragraph_format.first_line_indent = Cm(indent)
        if keep:
            par.paragraph_format.keep_with_next = True
        return par

    def bullets(self, items, numbered=False):
        for it in items:
            par = self.doc.add_paragraph(
                it, style='List Number' if numbered else 'List Bullet')
            pf = par.paragraph_format
            pf.line_spacing = 1.5
            pf.space_after = Pt(0)
            pf.left_indent = Cm(1.25)
            pf.first_line_indent = Cm(-0.5)
            for r in par.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(14)

    def code(self, text, size=9):
        for line in text.split('\n'):
            par = self.doc.add_paragraph(style='CodeBlock')
            r = par.add_run(line.replace('\t', '    '))
            r.font.size = Pt(size)

    def code_file(self, path, size=9):
        self.code(open(path, encoding='utf-8').read().rstrip('\n'), size)

    def caption(self, text):
        par = self.p(text, align=R, indent=0, keep=True)
        par.runs[0].font.size = Pt(13)
        par.paragraph_format.space_after = Pt(6)
        return par

    def table(self, headers, rows, widths=None, font=11, first_left=False):
        t = self.doc.add_table(rows=0, cols=len(headers))
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER

        def fill(cells, values, bold):
            for k, (cell, txt) in enumerate(zip(cells, values)):
                cell.text = ''
                par = cell.paragraphs[0]
                par.alignment = L if (first_left and k == 0 and not bold) else C
                par.paragraph_format.first_line_indent = Cm(0)
                par.paragraph_format.line_spacing = 1.0
                par.paragraph_format.space_after = Pt(0)
                r = par.add_run(str(txt))
                r.bold = bold
                r.font.name = 'Times New Roman'
                r.font.size = Pt(font)

        fill(t.add_row().cells, headers, True)
        for row in rows:
            fill(t.add_row().cells, row, False)
        if widths:
            for row in t.rows:
                for cell, w in zip(row.cells, widths):
                    cell.width = Cm(w)
        self.doc.add_paragraph()
        return t

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self, path):
        self.doc.save(path)
        return path
